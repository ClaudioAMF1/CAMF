"""Geração de relatórios PDF (WeasyPrint), XLSX (openpyxl) e CSV."""

import csv
import io
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
from jinja2 import Environment, FileSystemLoader
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import settings

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"
FORMATO_BRL = 'R$ #,##0.00'


def formatar_brl(valor) -> str:
    if valor is None:
        return ""
    num = f"{Decimal(str(valor)):,.2f}"
    return "R$ " + num.replace(",", "_").replace(".", ",").replace("_", ".")


def formatar_data(valor) -> str:
    return f"{valor:%d/%m/%Y}" if valor else ""


def gerar_pdf(df: pd.DataFrame, resumo: dict, periodo: str) -> bytes:
    from weasyprint import HTML

    env = Environment(loader=FileSystemLoader(TEMPLATES_DIR))
    env.filters["brl"] = formatar_brl
    env.filters["data_br"] = formatar_data
    template = env.get_template("relatorio.html")

    detalhado = df.sort_values(["pagador_nome", "vencimento"]).to_dict("records") if not df.empty else []
    html = template.render(
        empresa=settings.empresa_nome,
        cnpj=settings.empresa_cnpj,
        gerado_em=datetime.now(),
        periodo=periodo,
        resumo=resumo,
        pagadores=resumo_por_pagador(df),
        detalhado=detalhado,
    )
    return HTML(string=html).write_pdf()


COLUNAS_DETALHADO = [
    ("ID", "id", 8),
    ("Pagador", "pagador_nome", 38),
    ("CPF/CNPJ", "pagador_cpf_cnpj", 20),
    ("Município", "pagador_municipio", 20),
    ("UF", "pagador_uf", 6),
    ("Nº Documento", "num_documento", 14),
    ("Nosso Número", "nosso_numero", 14),
    ("Emissão", "data_emissao", 12),
    ("Vencimento", "vencimento", 12),
    ("Valor", "valor", 14),
    ("Situação", "situacao", 12),
    ("Qualidade", "qualidade", 16),
    ("Vencido", "vencido", 10),
    ("Data Pagamento", "data_pagamento", 14),
    ("Valor Pago", "valor_pago", 14),
    ("Divergências", "divergencias", 30),
    ("Linha Digitável", "linha_digitavel", 50),
    ("Observação", "observacao", 30),
]

COLUNAS_PAGADORES = [
    ("Pagador", "nome", 38),
    ("CPF/CNPJ", "cpf_cnpj", 20),
    ("Endereço", "endereco", 34),
    ("Bairro", "bairro", 24),
    ("Município", "municipio", 20),
    ("UF", "uf", 6),
    ("CEP", "cep", 12),
    ("Qtd. boletos", "qtd", 13),
    ("Total", "total", 15),
    ("Em aberto", "total_aberto", 15),
    ("Pago", "total_pago", 15),
    ("Vencido", "total_vencido", 15),
]


def resumo_por_pagador(df: pd.DataFrame) -> list[dict]:
    """Uma linha por pessoa, com os dados cadastrais e os totais dela."""
    if df.empty:
        return []
    linhas = []
    for (nome, cpf), grupo in df.groupby(["pagador_nome", "pagador_cpf_cnpj"], dropna=False):
        primeiro = grupo.iloc[0]
        linhas.append({
            "nome": nome,
            "cpf_cnpj": cpf or "",
            "endereco": primeiro.get("pagador_endereco") or "",
            "bairro": primeiro.get("pagador_bairro") or "",
            "municipio": primeiro.get("pagador_municipio") or "",
            "uf": primeiro.get("pagador_uf") or "",
            "cep": formatar_cep(primeiro.get("pagador_cep")),
            "qtd": int(len(grupo)),
            "total": Decimal(str(grupo["valor"].sum())),
            "total_aberto": Decimal(str(grupo.loc[grupo["situacao"] == "aberto", "valor"].sum())),
            "total_pago": Decimal(str(grupo.loc[grupo["situacao"] == "pago", "valor"].sum())),
            "total_vencido": Decimal(str(grupo.loc[grupo["vencido"], "valor"].sum())),
        })
    return sorted(linhas, key=lambda x: x["nome"])


def formatar_cep(cep) -> str:
    digitos = "".join(ch for ch in str(cep or "") if ch.isdigit())
    return f"{digitos[:5]}-{digitos[5:]}" if len(digitos) == 8 else ""


def gerar_xlsx(df: pd.DataFrame, resumo: dict) -> bytes:
    wb = Workbook()
    negrito = Font(bold=True)
    cabecalho_fill = PatternFill("solid", fgColor="1F4E78")
    cabecalho_font = Font(bold=True, color="FFFFFF")

    # --- Aba Resumo: pagador × valor unitário (com o documento de cada um) ---
    docs = {}
    if not df.empty:
        docs = (
            df.dropna(subset=["pagador_nome"])
            .groupby("pagador_nome")["pagador_cpf_cnpj"]
            .first()
            .to_dict()
        )

    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Pagador", "CPF/CNPJ", "Qtd", "Valor Unitário", "Subtotal"])
    for col in range(1, 6):
        celula = ws.cell(row=1, column=col)
        celula.fill = cabecalho_fill
        celula.font = cabecalho_font
    for item in resumo["por_pagador"]:
        ws.append([
            item["nome"], docs.get(item["nome"]) or "",
            item["qtd"], item["valor_unitario"], item["subtotal"],
        ])
        ws.cell(row=ws.max_row, column=4).number_format = FORMATO_BRL
        ws.cell(row=ws.max_row, column=5).number_format = FORMATO_BRL
    ws.append(["TOTAL GERAL", None, resumo["qtd_boletos"], None, resumo["total_geral"]])
    for col in range(1, 6):
        ws.cell(row=ws.max_row, column=col).font = negrito
    ws.cell(row=ws.max_row, column=5).number_format = FORMATO_BRL
    for col, largura in zip(range(1, 6), (40, 20, 8, 16, 16)):
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.freeze_panes = "A2"

    # --- Aba Pagadores: cadastro + totais por pessoa ---
    ws_pag = wb.create_sheet("Pagadores")
    ws_pag.append([titulo for titulo, _, _ in COLUNAS_PAGADORES])
    for col in range(1, len(COLUNAS_PAGADORES) + 1):
        celula = ws_pag.cell(row=1, column=col)
        celula.fill = cabecalho_fill
        celula.font = cabecalho_font
        celula.alignment = Alignment(horizontal="center")
    for pessoa in resumo_por_pagador(df):
        ws_pag.append([pessoa[chave] for _, chave, _ in COLUNAS_PAGADORES])
        for idx, (_, chave, _) in enumerate(COLUNAS_PAGADORES, start=1):
            if chave.startswith("total"):
                ws_pag.cell(row=ws_pag.max_row, column=idx).number_format = FORMATO_BRL
    for idx, (_, _, largura) in enumerate(COLUNAS_PAGADORES, start=1):
        ws_pag.column_dimensions[get_column_letter(idx)].width = largura
    ws_pag.freeze_panes = "A2"

    # --- Aba Detalhado ---
    ws2 = wb.create_sheet("Detalhado")
    ws2.append([titulo for titulo, _, _ in COLUNAS_DETALHADO])
    for col in range(1, len(COLUNAS_DETALHADO) + 1):
        celula = ws2.cell(row=1, column=col)
        celula.fill = cabecalho_fill
        celula.font = cabecalho_font
        celula.alignment = Alignment(horizontal="center")
    registros = df.to_dict("records") if not df.empty else []
    for registro in registros:
        linha = []
        for _, chave, _ in COLUNAS_DETALHADO:
            valor = registro.get(chave)
            if chave in ("data_emissao", "vencimento", "data_pagamento"):
                valor = formatar_data(valor)
            elif chave == "vencido":
                valor = "sim" if valor else "não"
            elif valor is not None and not isinstance(valor, (str, int, float, Decimal)):
                valor = str(valor)
            linha.append(valor)
        ws2.append(linha)
        for idx, (_, chave, _) in enumerate(COLUNAS_DETALHADO, start=1):
            if chave in ("valor", "valor_pago"):
                ws2.cell(row=ws2.max_row, column=idx).number_format = FORMATO_BRL
    for idx, (_, _, largura) in enumerate(COLUNAS_DETALHADO, start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = largura
    ws2.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_csv_pagadores(df: pd.DataFrame) -> bytes:
    """CSV com uma linha por pessoa (cadastro + totais)."""
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";")
    escritor.writerow([titulo for titulo, _, _ in COLUNAS_PAGADORES])
    for pessoa in resumo_por_pagador(df):
        linha = []
        for _, chave, _ in COLUNAS_PAGADORES:
            valor = pessoa[chave]
            if chave.startswith("total"):
                valor = f"{valor:.2f}".replace(".", ",")
            linha.append(valor)
        escritor.writerow(linha)
    return buffer.getvalue().encode("utf-8-sig")


def gerar_csv(df: pd.DataFrame) -> bytes:
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";")
    escritor.writerow([titulo for titulo, _, _ in COLUNAS_DETALHADO])
    registros = df.to_dict("records") if not df.empty else []
    for registro in registros:
        linha = []
        for _, chave, _ in COLUNAS_DETALHADO:
            valor = registro.get(chave)
            if chave in ("data_emissao", "vencimento", "data_pagamento"):
                valor = formatar_data(valor)
            elif chave in ("valor", "valor_pago"):
                valor = str(valor).replace(".", ",") if valor is not None else ""
            elif chave == "vencido":
                valor = "sim" if valor else "não"
            linha.append(valor if valor is not None else "")
        escritor.writerow(linha)
    return buffer.getvalue().encode("utf-8-sig")
