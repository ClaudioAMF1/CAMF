from datetime import date
from decimal import Decimal

from .conftest import enviar, texto_boleto

CPF_A = "529.982.247-25"
CNPJ_CAMF = "42.800.118/0001-44"
VENC = date(2026, 8, 10)


def _pagina(seq=1, nome="JOAO DA SILVA", cpf=CPF_A, valor=Decimal("100.00"), **kw):
    return texto_boleto(nome, cpf, valor, VENC, sequencia=seq, **kw)


# ---------- Upload e extração ----------

def test_upload_cria_boleto(client):
    resp = enviar(client, {"a.pdf": [_pagina(1)]})
    assert resp.status_code == 200
    arquivo = resp.json()["arquivos"][0]
    assert arquivo["novos"] == 1
    assert arquivo["duplicados"] == 0
    assert arquivo["erro"] is None

    boletos = client.get("/api/boletos").json()
    assert boletos["total"] == 1
    boleto = boletos["items"][0]
    assert boleto["valor"] == "100.00"
    assert boleto["vencimento"] == "2026-08-10"
    assert boleto["situacao"] == "aberto"
    assert boleto["qualidade"] == "ok"
    assert boleto["pagador_nome"] == "JOAO DA SILVA"
    assert boleto["vencido"] is False


def test_pagina_sem_linha_digitavel_ignorada(client):
    resp = enviar(client, {"a.pdf": [_pagina(1), "Página de carta sem boleto nenhum"]})
    arquivo = resp.json()["arquivos"][0]
    assert arquivo["novos"] == 1
    assert arquivo["ignoradas"] == [2]

    upload = client.get(f"/api/uploads/{arquivo['upload_id']}").json()
    assert upload["qtd_ignoradas"] == 1
    assert upload["qtd_paginas"] == 2


def test_divergencia_valor_marca_revisao(client):
    pagina = _pagina(1, valor=Decimal("100.00"), valor_texto=Decimal("999.00"))
    resp = enviar(client, {"a.pdf": [pagina]})
    assert resp.json()["arquivos"][0]["revisao_manual"] == 1

    boleto = client.get("/api/boletos").json()["items"][0]
    assert boleto["qualidade"] == "revisao_manual"
    assert "valor_divergente" in boleto["divergencias"]
    # linha digitável é autoritativa; valor do texto vai para a observação
    assert boleto["valor"] == "100.00"
    assert "999" in boleto["observacao"]


def test_cpf_invalido_gera_pagador_provisorio(client):
    pagina = _pagina(1, cpf="111.111.111-11")
    enviar(client, {"a.pdf": [pagina]})
    boleto = client.get("/api/boletos").json()["items"][0]
    assert boleto["qualidade"] == "revisao_manual"
    assert "cpf_cnpj_ausente" in boleto["divergencias"]

    pagador = client.get(f"/api/pagadores/{boleto['pagador_id']}").json()
    assert pagador["provisorio"] is True
    assert pagador["cpf_cnpj"] is None


def test_erro_por_arquivo_nao_derruba_lote(client):
    files = [
        ("arquivos", ("bom.pdf", b"%PDF" + _pagina(1).encode(), "application/pdf")),
        ("arquivos", ("ruim.pdf", b"nao e pdf", "application/pdf")),
    ]
    resp = client.post("/api/uploads", files=files)
    assert resp.status_code == 200
    por_nome = {a["nome"]: a for a in resp.json()["arquivos"]}
    assert por_nome["bom.pdf"]["novos"] == 1
    assert por_nome["ruim.pdf"]["erro"] is not None


# ---------- Deduplicação e reprocessamento ----------

def test_dedup_por_linha_digitavel(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    # arquivo diferente (hash diferente), mesmo boleto
    resp = enviar(client, {"b.pdf": [_pagina(1), _pagina(2)]})
    arquivo = resp.json()["arquivos"][0]
    assert arquivo["duplicados"] == 1
    assert arquivo["novos"] == 1
    assert client.get("/api/boletos").json()["total"] == 2


def test_arquivo_repetido_retorna_409(client):
    resp1 = enviar(client, {"a.pdf": [_pagina(1)]})
    upload_id = resp1.json()["arquivos"][0]["upload_id"]

    resp2 = enviar(client, {"a.pdf": [_pagina(1)]})
    assert resp2.status_code == 409
    detalhe = resp2.json()["detail"]
    assert detalhe["uploads_originais"][0]["upload"]["id"] == upload_id


def test_reprocessamento_com_forcar(client):
    resp1 = enviar(client, {"a.pdf": [_pagina(1)]})
    original = resp1.json()["arquivos"][0]["upload_id"]

    resp2 = enviar(client, {"a.pdf": [_pagina(1)]}, forcar=True)
    assert resp2.status_code == 200
    arquivo = resp2.json()["arquivos"][0]
    assert arquivo["upload_id"] != original
    # reprocessar = re-extrair: o boleto existente é atualizado, não ignorado
    assert arquivo["atualizados"] == 1
    assert arquivo["duplicados"] == 0
    assert arquivo["novos"] == 0

    novo_upload = client.get(f"/api/uploads/{arquivo['upload_id']}").json()
    assert novo_upload["reprocessado_de_id"] == original
    assert novo_upload["qtd_atualizados"] == 1
    # sem boleto duplicado: a linha digitável continua única
    assert client.get("/api/boletos").json()["total"] == 1


def test_reparo_reenviando_com_forcar_sem_deletar_nada(client):
    """Fluxo de reparo em um passo: corrigido o parser, basta reenviar o mesmo
    PDF com forcar=true — os boletos existentes são re-extraídos no lugar."""
    # extração ruim: sem CPF, nome capturado errado -> pagador provisório
    resp = enviar(client, {"a.pdf": [_pagina(1, nome="Nome do pagador Numero", cpf=None)]})
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]
    assert resp.json()["arquivos"][0]["revisao_manual"] == 1

    # parser corrigido: mesmo arquivo reenviado, nome e CPF agora certos
    resp2 = enviar(client, {"a.pdf": [_pagina(1, nome="JOAO DA SILVA", cpf=CPF_A)]}, forcar=True)
    arquivo = resp2.json()["arquivos"][0]
    assert arquivo["atualizados"] == 1
    assert arquivo["duplicados"] == 0

    boletos = client.get("/api/boletos").json()
    assert boletos["total"] == 1
    boleto = boletos["items"][0]
    assert boleto["id"] == boleto_id  # mesmo registro, re-extraído
    assert boleto["pagador_nome"] == "JOAO DA SILVA"
    assert boleto["qualidade"] == "ok"
    assert boleto["divergencias"] == []

    # a trilha de auditoria do registro é preservada e ganha a re-extração
    trilha = client.get(f"/api/boletos/{boleto_id}/auditoria").json()
    assert any(r["campo"] == "reextracao" for r in trilha)
    assert any(r["acao"] == "criar" for r in trilha)


def test_reprocessamento_preserva_pagamento_ja_registrado(client):
    """Re-extrair não pode apagar o que o usuário registrou: situação e
    pagamento sobrevivem ao reprocessamento."""
    enviar(client, {"a.pdf": [_pagina(1)]})
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]
    client.post(
        f"/api/boletos/{boleto_id}/pagar",
        json={"data_pagamento": "2026-08-01", "valor_pago": "100.00"},
    )

    enviar(client, {"a.pdf": [_pagina(1)]}, forcar=True)

    boleto = client.get(f"/api/boletos/{boleto_id}").json()
    assert boleto["situacao"] == "pago"
    assert boleto["data_pagamento"] == "2026-08-01"
    assert boleto["valor_pago"] == "100.00"


def test_reparo_apos_deletar_upload_tambem_funciona(client):
    """O caminho antigo (deletar upload + reenviar) continua válido."""
    resp = enviar(client, {"a.pdf": [_pagina(1, nome="Nome do pagador Numero", cpf=None)]})
    upload_id = resp.json()["arquivos"][0]["upload_id"]
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]

    client.delete(f"/api/uploads/{upload_id}")
    resp2 = enviar(client, {"a2.pdf": [_pagina(1, nome="JOAO DA SILVA", cpf=CPF_A)]}, forcar=True)

    assert resp2.json()["arquivos"][0]["atualizados"] == 1
    boletos = client.get("/api/boletos").json()
    assert boletos["total"] == 1
    assert boletos["items"][0]["id"] == boleto_id  # restaurado e re-extraído
    assert boletos["items"][0]["pagador_nome"] == "JOAO DA SILVA"

    # o pagador provisório fica órfão e pode ser removido
    orfao = next(p for p in client.get("/api/pagadores").json() if p["provisorio"])
    assert orfao["qtd_boletos"] == 0
    assert client.delete(f"/api/pagadores/{orfao['id']}").status_code == 200
    assert all(p["provisorio"] is False for p in client.get("/api/pagadores").json())


def test_delete_pagador_com_boletos_ativos_recusado(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    pagador_id = client.get("/api/pagadores").json()[0]["id"]
    assert client.delete(f"/api/pagadores/{pagador_id}").status_code == 400


def test_diagnostico_nao_grava(client):
    from .conftest import fake_pdf

    files = [("arquivos", ("a.pdf", fake_pdf([_pagina(1)]), "application/pdf"))]
    resp = client.post("/api/uploads/analisar", files=files)
    assert resp.status_code == 200
    pagina = resp.json()["arquivos"][0]["paginas"][0]
    assert pagina["tem_linha_digitavel"] is True
    assert pagina["campos"]["pagador_nome"] == "JOAO DA SILVA"
    assert pagina["divergencias"] == []
    # nada foi gravado
    assert client.get("/api/boletos").json()["total"] == 0
    assert client.get("/api/uploads").json() == []


# ---------- Soft delete ----------

def test_soft_delete_boleto_some_das_queries(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]

    resp = client.delete(f"/api/boletos/{boleto_id}")
    assert resp.status_code == 200
    assert resp.json()["deletado_em"] is not None

    assert client.get("/api/boletos").json()["total"] == 0
    assert client.get("/api/boletos?incluir_deletados=true").json()["total"] == 1

    client.post(f"/api/boletos/{boleto_id}/restaurar")
    assert client.get("/api/boletos").json()["total"] == 1


def test_soft_delete_upload_cascata_e_restauracao(client):
    resp = enviar(client, {"a.pdf": [_pagina(1), _pagina(2)]})
    upload_id = resp.json()["arquivos"][0]["upload_id"]

    client.delete(f"/api/uploads/{upload_id}")
    assert client.get("/api/boletos").json()["total"] == 0
    assert all(u["id"] != upload_id for u in client.get("/api/uploads").json())

    client.post(f"/api/uploads/{upload_id}/restaurar")
    assert client.get("/api/boletos").json()["total"] == 2


# ---------- Pagamento, PATCH e vencido derivado ----------

def test_marcar_pago(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]

    resp = client.post(
        f"/api/boletos/{boleto_id}/pagar",
        json={"data_pagamento": "2026-08-01", "valor_pago": "100.00"},
    )
    corpo = resp.json()
    assert corpo["situacao"] == "pago"
    assert corpo["data_pagamento"] == "2026-08-01"
    assert corpo["vencido"] is False  # pago nunca é vencido


def test_vencido_derivado_e_filtro(client):
    vencido = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("50.00"), date(2026, 1, 10), sequencia=1)
    futuro = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("60.00"), date(2030, 1, 10), sequencia=2)
    enviar(client, {"a.pdf": [vencido, futuro]})

    todos = client.get("/api/boletos").json()["items"]
    assert {b["vencido"] for b in todos} == {True, False}

    filtrados = client.get("/api/boletos?vencido=true").json()
    assert filtrados["total"] == 1
    assert filtrados["items"][0]["valor"] == "50.00"


def test_patch_revalida_qualidade(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]

    # alterar o valor para algo diferente da linha digitável -> revisão manual
    resp = client.patch(f"/api/boletos/{boleto_id}", json={"valor": "123.45"})
    assert resp.json()["qualidade"] == "revisao_manual"
    assert "valor_divergente" in resp.json()["divergencias"]

    # voltar ao valor da linha -> qualidade ok de novo
    resp = client.patch(f"/api/boletos/{boleto_id}", json={"valor": "100.00"})
    assert resp.json()["qualidade"] == "ok"
    assert resp.json()["divergencias"] == []


# ---------- Auditoria ----------

def test_auditoria_criacao_e_edicao(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]

    trilha = client.get(f"/api/boletos/{boleto_id}/auditoria").json()
    assert any(r["acao"] == "criar" and r["origem"] == "extracao" for r in trilha)

    client.patch(
        f"/api/boletos/{boleto_id}",
        json={"observacao": "conferido"},
        headers={"X-Autor": "maria"},
    )
    trilha = client.get(f"/api/boletos/{boleto_id}/auditoria").json()
    edicao = next(r for r in trilha if r["acao"] == "editar" and r["campo"] == "observacao")
    assert edicao["valor_novo"] == "conferido"
    assert edicao["autor"] == "maria"
    assert edicao["origem"] == "manual"


def test_auditoria_marcar_pago_e_delete(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    boleto_id = client.get("/api/boletos").json()["items"][0]["id"]
    client.post(
        f"/api/boletos/{boleto_id}/pagar",
        json={"data_pagamento": "2026-08-01", "valor_pago": "100.00"},
    )
    client.delete(f"/api/boletos/{boleto_id}")
    client.post(f"/api/boletos/{boleto_id}/restaurar")

    acoes = [r["acao"] for r in client.get(f"/api/boletos/{boleto_id}/auditoria").json()]
    for acao in ("criar", "marcar_pago", "deletar", "restaurar"):
        assert acao in acoes


# ---------- Identidade do pagador ----------

def test_nome_divergente_nao_sobrescreve_e_vira_alternativo(client):
    enviar(client, {"a.pdf": [_pagina(1, nome="JOAO DA SILVA")]})
    enviar(client, {"b.pdf": [_pagina(2, nome="JOAO D SILVA")]})

    pagadores = client.get("/api/pagadores").json()
    assert len(pagadores) == 1
    pagador = pagadores[0]
    assert pagador["nome"] == "JOAO DA SILVA"
    assert "JOAO D SILVA" in pagador["nomes_alternativos"]
    assert pagador["qtd_boletos"] == 2


def test_nomes_iguais_cpfs_diferentes_sao_pagadores_distintos(client):
    outro_cpf = "11.222.333/0001-81"
    enviar(client, {"a.pdf": [_pagina(1, nome="JOAO DA SILVA", cpf=CPF_A)]})
    enviar(client, {"b.pdf": [_pagina(2, nome="JOAO DA SILVA", cpf=outro_cpf)]})

    assert len(client.get("/api/pagadores").json()) == 2
    sugestoes = client.get("/api/pagadores/sugestoes-merge").json()
    assert len(sugestoes) == 1
    assert len(sugestoes[0]["pagadores"]) == 2


def test_merge_pagadores(client):
    enviar(client, {"a.pdf": [_pagina(1, cpf=CPF_A)]})
    enviar(client, {"b.pdf": [_pagina(2, cpf="11.222.333/0001-81")]})
    pagadores = client.get("/api/pagadores").json()
    destino, origem = pagadores[0], pagadores[1]

    resp = client.post(
        f"/api/pagadores/{destino['id']}/merge",
        json={"pagador_origem_id": origem["id"]},
    )
    assert resp.status_code == 200
    assert resp.json()["qtd_boletos"] == 2

    restantes = client.get("/api/pagadores").json()
    assert [p["id"] for p in restantes] == [destino["id"]]
    boletos = client.get("/api/boletos").json()["items"]
    assert all(b["pagador_id"] == destino["id"] for b in boletos)


def test_correcao_cpf_provisorio_faz_merge(client):
    enviar(client, {"a.pdf": [_pagina(1, cpf=CPF_A)]})  # pagador real
    enviar(client, {"b.pdf": [_pagina(2, cpf=None)]})   # provisório, mesmo nome

    pagadores = client.get("/api/pagadores").json()
    real = next(p for p in pagadores if not p["provisorio"])
    provisorio = next(p for p in pagadores if p["provisorio"])

    resp = client.patch(
        f"/api/pagadores/{provisorio['id']}", json={"cpf_cnpj": CPF_A}
    )
    assert resp.status_code == 200
    assert resp.json()["id"] == real["id"]
    assert resp.json()["qtd_boletos"] == 2
    assert len(client.get("/api/pagadores").json()) == 1


# ---------- Agrupamento e baixa em lote ----------

def test_agrupamento_por_pagador(client):
    outro = "11.222.333/0001-81"
    enviar(client, {"a.pdf": [
        _pagina(1, nome="JOAO DA SILVA", cpf=CPF_A, valor=Decimal("100.00")),
        _pagina(2, nome="JOAO DA SILVA", cpf=CPF_A, valor=Decimal("200.00")),
        _pagina(3, nome="EMPRESA XYZ", cpf=outro, valor=Decimal("50.00")),
    ]})
    grupos = client.get("/api/boletos/por-pagador").json()
    assert len(grupos) == 2  # uma linha por pessoa, nome não se repete

    joao = next(g for g in grupos if g["nome"] == "JOAO DA SILVA")
    assert joao["qtd"] == 2
    assert joao["total"] == "300.00"
    assert joao["qtd_aberto"] == 2
    assert joao["total_aberto"] == "300.00"
    assert joao["qtd_pago"] == 0
    assert joao["proximo_vencimento"] == "2026-08-10"
    # ordenado por total desc
    assert grupos[0]["nome"] == "JOAO DA SILVA"

    # ao expandir: os boletos daquela pessoa
    boletos = client.get(f"/api/boletos?pagador_id={joao['pagador_id']}").json()
    assert boletos["total"] == 2


def test_agrupamento_respeita_filtros(client):
    enviar(client, {"a.pdf": [
        _pagina(1, valor=Decimal("100.00")),
        _pagina(2, valor=Decimal("900.00")),
    ]})
    grupos = client.get("/api/boletos/por-pagador?valor_min=500").json()
    assert len(grupos) == 1
    assert grupos[0]["qtd"] == 1
    assert grupos[0]["total"] == "900.00"


def test_agrupamento_conta_pagos_e_vencidos(client):
    vencido = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("70.00"), date(2020, 1, 10), sequencia=1)
    futuro = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("30.00"), date(2030, 1, 10), sequencia=2)
    enviar(client, {"a.pdf": [vencido, futuro]})

    alvo = client.get("/api/boletos?valor_min=30&valor_max=30").json()["items"][0]
    client.post(f"/api/boletos/{alvo['id']}/pagar",
                json={"data_pagamento": "2026-01-05", "valor_pago": "30.00"})

    grupo = client.get("/api/boletos/por-pagador").json()[0]
    assert grupo["qtd_pago"] == 1 and grupo["total_pago"] == "30.00"
    assert grupo["qtd_vencido"] == 1 and grupo["total_vencido"] == "70.00"


def test_pagar_lote(client):
    enviar(client, {"a.pdf": [
        _pagina(1, valor=Decimal("100.00")),
        _pagina(2, valor=Decimal("250.00")),
    ]})
    ids = [b["id"] for b in client.get("/api/boletos").json()["items"]]

    resp = client.post("/api/boletos/pagar-lote",
                       json={"ids": ids, "data_pagamento": "2026-08-01"})
    assert resp.json()["pagos"] == 2

    boletos = client.get("/api/boletos").json()["items"]
    assert all(b["situacao"] == "pago" for b in boletos)
    # sem valor_pago informado, cada boleto é baixado pelo próprio valor
    assert {b["valor_pago"] for b in boletos} == {"100.00", "250.00"}

    # reenviar o mesmo lote não dá erro: já pagos são ignorados
    resp2 = client.post("/api/boletos/pagar-lote",
                        json={"ids": ids, "data_pagamento": "2026-08-01"})
    assert resp2.json()["pagos"] == 0
    assert sorted(resp2.json()["ignorados"]) == sorted(ids)


# ---------- Dashboard e relatórios ----------

def test_dashboard_agregacoes(client):
    enviar(client, {"a.pdf": [
        _pagina(1, valor=Decimal("100.00")),
        _pagina(2, valor=Decimal("100.00")),
        _pagina(3, valor=Decimal("250.00")),
    ]})
    dash = client.get("/api/dashboard").json()
    assert dash["qtd_boletos"] == 3
    assert dash["total_geral"] == "450.00"
    assert dash["qtd_pagadores"] == 1
    grupos = {(g["valor_unitario"], g["qtd"]) for g in dash["por_pagador"]}
    assert grupos == {("100.00", 2), ("250.00", 1)}


def test_relatorio_csv_e_xlsx(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    csv_resp = client.get("/api/relatorios/csv")
    assert csv_resp.status_code == 200
    assert "JOAO DA SILVA" in csv_resp.text

    xlsx_resp = client.get("/api/relatorios/xlsx")
    assert xlsx_resp.status_code == 200
    assert xlsx_resp.content[:2] == b"PK"  # zip/xlsx


# ---------- Vencimentos e exportações com dados do pagador ----------

def test_alertas_aceita_intervalo_do_mes(client):
    dentro = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("10.00"), date(2026, 8, 20), sequencia=1)
    fora = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("20.00"), date(2026, 9, 20), sequencia=2)
    enviar(client, {"a.pdf": [dentro, fora]})

    # todos os vencimentos de agosto/2026
    alertas = client.get("/api/dashboard/alertas?de=2026-08-01&ate=2026-08-31").json()
    assert len(alertas) == 1
    assert alertas[0]["valor"] == "10.00"
    assert alertas[0]["pagador_cpf_cnpj"] == "529.982.247-25"


def test_alertas_pode_incluir_vencidos(client):
    vencido = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("15.00"), date(2020, 3, 10), sequencia=1)
    futuro = texto_boleto("JOAO DA SILVA", CPF_A, Decimal("25.00"), date(2030, 3, 10), sequencia=2)
    enviar(client, {"a.pdf": [vencido, futuro]})

    sem = client.get("/api/dashboard/alertas?dias=365").json()
    assert all(not b["vencido"] for b in sem)

    com = client.get("/api/dashboard/alertas?ate=2030-12-31&incluir_vencidos=true").json()
    assert any(b["vencido"] for b in com)


def test_csv_de_pagadores_traz_cadastro_e_totais(client):
    enviar(client, {"a.pdf": [
        _pagina(1, nome="JOAO DA SILVA", cpf=CPF_A, valor=Decimal("100.00")),
        _pagina(2, nome="JOAO DA SILVA", cpf=CPF_A, valor=Decimal("50.00")),
    ]})
    texto = client.get("/api/relatorios/csv?aba=pagadores").text
    assert "CPF/CNPJ" in texto and "Qtd. boletos" in texto
    assert "529.982.247-25" in texto  # documento formatado
    assert "JOAO DA SILVA" in texto
    linhas = [l for l in texto.strip().splitlines() if l.strip()]
    assert len(linhas) == 2  # cabeçalho + uma pessoa
    assert "150,00" in linhas[1]  # total somado


def test_csv_de_boletos_traz_o_documento_do_pagador(client):
    enviar(client, {"a.pdf": [_pagina(1)]})
    texto = client.get("/api/relatorios/csv").text
    assert "529.982.247-25" in texto


def test_xlsx_tem_aba_de_pagadores(client):
    import io as _io
    from openpyxl import load_workbook

    enviar(client, {"a.pdf": [_pagina(1)]})
    resp = client.get("/api/relatorios/xlsx")
    wb = load_workbook(_io.BytesIO(resp.content))
    assert wb.sheetnames == ["Resumo", "Pagadores", "Detalhado"]

    aba = wb["Pagadores"]
    cabecalhos = [c.value for c in aba[1]]
    assert "CPF/CNPJ" in cabecalhos and "Município" in cabecalhos
    assert aba.cell(row=2, column=2).value == "529.982.247-25"

    # o documento também aparece no Resumo
    assert "CPF/CNPJ" in [c.value for c in wb["Resumo"][1]]
