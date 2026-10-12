"""Exportações no layout dos relatórios impressos (parcelas em aberto e títulos)."""

import io
from datetime import date, timedelta
from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from app.services import modelos
from app.services.linha_digitavel import montar_linha

from .conftest import enviar, texto_boleto

CPF_A = "529.982.247-25"
CPF_B = "111.444.777-35"
DIA_1 = date.today() + timedelta(days=10)
DIA_2 = date.today() + timedelta(days=20)


def _carregar(client):
    """Três abertos em dois dias e um pago, que não pode aparecer."""
    paginas = [
        texto_boleto("MARIA SOUZA", CPF_A, Decimal("1000.00"), DIA_1, sequencia=1,
                     num_documento="900-02", nosso_numero="5-1"),
        texto_boleto("JOAO LIMA", CPF_B, Decimal("600.00"), DIA_1, sequencia=2,
                     num_documento="51-03", nosso_numero="3-2"),
        texto_boleto("MARIA SOUZA", CPF_A, Decimal("2500.50"), DIA_2, sequencia=3,
                     num_documento="900-03", nosso_numero="7-3"),
        texto_boleto("JOAO LIMA", CPF_B, Decimal("99.00"), DIA_2, sequencia=4,
                     num_documento="99-01", nosso_numero="9-9"),
    ]
    enviar(client, {"a.pdf": paginas})
    pago = next(b for b in client.get("/api/boletos").json()["items"] if b["num_documento"] == "99-01")
    client.post(f"/api/boletos/{pago['id']}/pagar", json={"data_pagamento": str(date.today()), "valor_pago": "99.00"})


def _planilha(resp):
    return load_workbook(io.BytesIO(resp.content)).active


def _linhas(ws):
    return [[c for c in linha] for linha in ws.iter_rows(values_only=True)]


# ---------- Cedente ----------

def test_campo_livre_sicoob_do_boleto_real():
    """Linha digitável real: cooperativa 5004, cedente 185879-3."""
    info = modelos.campo_livre_sicoob("75691.50043 01185.879309 00000.940015 2 15750000960000")
    assert info == {"cooperativa": "5004", "cedente": "185879-3"}


def test_campo_livre_fora_do_padrao_nao_inventa_cedente():
    linha = montar_linha(campo_livre="0" * 24 + "1", valor=Decimal("10"), vencimento=DIA_1)
    assert modelos.campo_livre_sicoob(linha) is None
    assert modelos.campo_livre_sicoob(None) is None


def test_cedente_mais_frequente():
    sicoob = montar_linha(campo_livre="1500401185879300000094001",
                          valor=Decimal("10"), vencimento=DIA_1)
    df = pd.DataFrame({"linha_digitavel": [sicoob, sicoob, None]})
    cedente = modelos.identificar_cedente(df)
    assert (cedente.codigo, cedente.cooperativa) == ("185879-3", "5004")


# ---------- Parcelas em aberto ----------

def test_parcelas_agrupa_por_vencimento_e_ignora_pagos(client):
    _carregar(client)
    resp = client.get("/api/relatorios/parcelas-em-aberto.xlsx")
    assert resp.status_code == 200
    assert "parcelas_em_aberto_" in resp.headers["content-disposition"]

    linhas = _linhas(_planilha(resp))
    textos = [str(c) for linha in linhas for c in linha if c is not None]
    assert "99-01" not in textos
    assert textos.count("Total Cedente:") == 2
    assert textos.count("Total por Data de Vencimento:") == 2

    # dentro do dia, a ordem é a do nosso número (3-2 antes de 5-1)
    sacados = [linha[1] for linha in linhas if linha[2] in ("3-2", "5-1", "7-3")]
    assert sacados == ["JOAO LIMA", "MARIA SOUZA", "MARIA SOUZA"]

    total_dia_1 = next(l for l in linhas if l[1] == "Total por Data de Vencimento:")
    assert total_dia_1[6] == 1600 and total_dia_1[7] == 0 and total_dia_1[8] == 1600
    geral = next(l for l in linhas if l[1] == "Total Geral (3 parcela(s)):")
    assert geral[8] == 4100.5


def test_parcelas_monta_saldo_e_totais():
    df = pd.DataFrame([
        {"vencimento": DIA_1, "valor": Decimal("10.00"), "pagador_nome": "A",
         "nosso_numero": "2-1", "num_documento": "1", "especie": None},
        {"vencimento": DIA_1, "valor": Decimal("5.50"), "pagador_nome": "B",
         "nosso_numero": "10-1", "num_documento": "2", "especie": "DM"},
    ])
    dados = modelos.montar_parcelas_em_aberto(df)
    grupo = dados["grupos"][0]
    assert [i["nosso_numero"] for i in grupo["itens"]] == ["2-1", "10-1"]  # numérico, não texto
    assert grupo["itens"][0]["docto"] == "DM"  # espécie ausente vira DM
    assert grupo["total_saldo"] == Decimal("15.50")
    assert dados["total_valor"] == Decimal("15.50")


def test_parcelas_pdf(client):
    _carregar(client)
    resp = client.get("/api/relatorios/parcelas-em-aberto.pdf")
    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF")


def test_parcelas_respeita_filtro_de_datas(client):
    _carregar(client)
    resp = client.get(f"/api/relatorios/parcelas-em-aberto.xlsx?vencimento_ate={DIA_1}")
    textos = [str(c) for linha in _linhas(_planilha(resp)) for c in linha if c is not None]
    assert textos.count("Total por Data de Vencimento:") == 1
    # sem "de" no filtro, o início do período é o primeiro vencimento dos dados
    assert f"Período: {DIA_1:%d/%m/%Y} a {DIA_1:%d/%m/%Y}" in textos


def test_formato_invalido_404(client):
    assert client.get("/api/relatorios/parcelas-em-aberto.csv").status_code in (404, 422)


# ---------- Relação de títulos ----------

def test_titulos_xlsx_com_taxa(client):
    _carregar(client)
    resp = client.get("/api/relatorios/titulos.xlsx?taxa=1.69")
    assert resp.status_code == 200
    ws = _planilha(resp)
    linhas = _linhas(ws)
    assert linhas[0] == ["Tipo", "DOC", "Valor nominal", "Valor Disponível",
                         "Vencimento", "Tx.Juros", "Pagador"]
    corpo = [l for l in linhas[1:] if l[0] == "Boleto"]
    assert [l[1] for l in corpo] == ["3-2", "5-1", "7-3"]  # por vencimento, depois nosso nº
    assert all(abs(l[5] - 0.0169) < 1e-9 for l in corpo)  # 1,69% como fração
    assert ws.cell(row=2, column=6).number_format == "0.0000 %"
    assert all(l[2] == l[3] for l in corpo)  # disponível = nominal
    total = next(l for l in linhas if l[0] == "Total")
    assert total[1] == "3 título(s)" and total[2] == 4100.5


def test_titulos_sem_taxa_deixa_coluna_vazia(client):
    _carregar(client)
    linhas = _linhas(_planilha(client.get("/api/relatorios/titulos.xlsx")))
    assert all(l[5] is None for l in linhas[1:] if l[0] == "Boleto")


def test_titulos_taxa_invalida(client):
    assert client.get("/api/relatorios/titulos.xlsx?taxa=-1").status_code == 422


def test_titulos_pdf(client):
    _carregar(client)
    resp = client.get("/api/relatorios/titulos.pdf?taxa=1.69")
    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF")


def test_formatacao():
    assert modelos.formatar_numero(Decimal("13771.66")) == "13.771,66"
    assert modelos.formatar_numero(0) == "0,00"
    assert modelos.formatar_taxa(Decimal("1.69")) == "1,6900 %"
    assert modelos.formatar_taxa(None) == ""
